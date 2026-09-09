import unittest
import os
import io
import shutil
import docx
from datetime import datetime, timezone
from app import create_app
from models import db, User, JobPosting, JobApplication, Employee, EmployeeDocument, DocumentTemplate
from services.offer_letter_service import (
    generate_offer_letter_docx,
    get_active_offer_letter_template,
    normalize_internship_duration,
    ensure_default_templates_initialized,
    OfferLetterTemplateNotFoundError
)


class TestRoleBasedOfferLetterSystem(unittest.TestCase):
    def setUp(self):
        """Set up isolated test Flask application and database."""
        self.app = create_app('testing')
        self.app.config['WTF_CSRF_ENABLED'] = False
        
        self.client = self.app.test_client()
        self.ctx = self.app.app_context()
        self.ctx.push()
        
        db.create_all()

        # Seed default templates
        ensure_default_templates_initialized()

        # Create Admin User
        self.admin = User(
            email='admin_test@antimatrix.co.in',
            name='Admin User',
            role='admin',
            is_active=True
        )
        self.admin.set_password('AdminSecurePass123!')
        db.session.add(self.admin)
        db.session.commit()

        # 1. AI & ML Job & Candidate
        self.job_aiml = JobPosting(
            title='AI & Machine Learning Engineering Intern',
            department='AI & ML',
            location='Remote',
            short_description='AI & ML Internship at Anti-Matrix',
            description='Detailed description for AI & ML internship.',
            duration='1_month',
            employment_type='Internship',
            is_active=True
        )
        # 2. Application Development Job & Candidate
        self.job_appdev = JobPosting(
            title='Mobile & Application Development Intern',
            department='Application Development',
            location='Remote',
            short_description='Application Development Internship at Anti-Matrix',
            description='Detailed description for Application Development internship.',
            duration='3_months',
            employment_type='Internship',
            is_active=True
        )
        # 3. Data Analytics Job & Candidate
        self.job_data = JobPosting(
            title='Data Analytics & Business Insights Intern',
            department='Data Analytics',
            location='Remote',
            short_description='Data Analytics Internship at Anti-Matrix',
            description='Detailed description for Data Analytics internship.',
            duration='1_month',
            employment_type='Internship',
            is_active=True
        )
        # 4. Full Stack Development Job & Candidate
        self.job_fullstack = JobPosting(
            title='Full Stack Web & Software Development Intern',
            department='Full Stack Development',
            location='Remote',
            short_description='Full Stack Development Internship at Anti-Matrix',
            description='Detailed description for Full Stack Development internship.',
            duration='3_months',
            employment_type='Internship',
            is_active=True
        )
        # 5. Unconfigured Domain Job
        self.job_unconfigured = JobPosting(
            title='Culinary Arts Specialist',
            department='Culinary Arts',
            location='Remote',
            short_description='Culinary Arts at Anti-Matrix',
            description='Detailed description for Culinary Arts.',
            duration='1_month',
            employment_type='Internship',
            is_active=True
        )

        db.session.add_all([self.job_aiml, self.job_appdev, self.job_data, self.job_fullstack, self.job_unconfigured])
        db.session.commit()

        # Candidate Applications
        self.app_aiml = JobApplication(
            job_id=self.job_aiml.id,
            application_code='AM-APP-AI001',
            full_name='Aarav Sharma',
            email='aarav.sharma@example.com',
            phone='+91 9876543210',
            college='IIT Bombay',
            duration='1_month',
            joining_date='01/10/2026',
            resume_filename='aarav_resume.pdf',
            resume_path='uploads/resumes/aarav_resume.pdf',
            payment_status='paid',
            application_status='shortlisted',
            status='Shortlisted'
        )
        self.app_appdev = JobApplication(
            job_id=self.job_appdev.id,
            application_code='AM-APP-APP002',
            full_name='Bhavna Patel',
            email='bhavna.patel@example.com',
            phone='+91 9876543211',
            college='BITS Pilani',
            duration='3_months',
            joining_date='15/10/2026',
            resume_filename='bhavna_resume.pdf',
            resume_path='uploads/resumes/bhavna_resume.pdf',
            payment_status='paid',
            application_status='shortlisted',
            status='Shortlisted'
        )
        self.app_data = JobApplication(
            job_id=self.job_data.id,
            application_code='AM-APP-DATA003',
            full_name='Chetan Verma',
            email='chetan.verma@example.com',
            phone='+91 9876543212',
            college='Delhi Technological University',
            duration='1_month',
            joining_date='01/11/2026',
            resume_filename='chetan_resume.pdf',
            resume_path='uploads/resumes/chetan_resume.pdf',
            payment_status='paid',
            application_status='shortlisted',
            status='Shortlisted'
        )
        self.app_fullstack = JobApplication(
            job_id=self.job_fullstack.id,
            application_code='AM-APP-FS004',
            full_name='Divya Nair',
            email='divya.nair@example.com',
            phone='+91 9876543213',
            college='NIT Trichy',
            duration='3_months',
            joining_date='01/10/2026',
            resume_filename='divya_resume.pdf',
            resume_path='uploads/resumes/divya_resume.pdf',
            payment_status='paid',
            application_status='shortlisted',
            status='Shortlisted'
        )
        self.app_unconfigured = JobApplication(
            job_id=self.job_unconfigured.id,
            application_code='AM-APP-CHEF005',
            full_name='Elena Gilbert',
            email='elena@example.com',
            phone='+91 9876543214',
            college='Culinary Academy',
            duration='1_month',
            joining_date='01/10/2026',
            resume_filename='elena_resume.pdf',
            resume_path='uploads/resumes/elena_resume.pdf',
            payment_status='paid',
            application_status='shortlisted',
            status='Shortlisted'
        )

        db.session.add_all([self.app_aiml, self.app_appdev, self.app_data, self.app_fullstack, self.app_unconfigured])
        db.session.commit()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.ctx.pop()

    def login_admin(self):
        with self.client.session_transaction() as sess:
            sess['_user_id'] = str(self.admin.id)
            sess['_fresh'] = True

    def test_01_scenario_aiml_1_month(self):
        """Scenario 1: Candidate in AI & ML with 1 Month duration uses AI & ML template."""
        emp_doc, output_path = generate_offer_letter_docx(self.app_aiml)
        self.assertTrue(os.path.exists(output_path))
        self.assertIsNotNone(emp_doc.template_id)

        # Inspect generated DOCX
        doc = docx.Document(output_path)
        full_text = "\n".join([p.text for p in doc.paragraphs])

        self.assertIn('Aarav Sharma', full_text)
        self.assertIn(self.app_aiml.formatted_code, full_text)
        self.assertIn('AI & ML Intern', full_text)
        self.assertIn('Artificial Intelligence and Machine Learning', full_text)
        self.assertIn('1 Month', full_text)
        self.assertIn('01/10/2026', full_text)
        self.assertNotIn('[Candidate Name]', full_text)
        self.assertNotIn('[1 Month / 3 Months]', full_text)
        self.assertNotIn('[Joining Date]', full_text)

    def test_02_scenario_app_dev_3_months(self):
        """Scenario 2: Candidate in Application Development with 3 Months duration uses App Dev template."""
        emp_doc, output_path = generate_offer_letter_docx(self.app_appdev)
        self.assertTrue(os.path.exists(output_path))

        doc = docx.Document(output_path)
        full_text = "\n".join([p.text for p in doc.paragraphs])

        self.assertIn('Bhavna Patel', full_text)
        self.assertIn(self.app_appdev.formatted_code, full_text)
        self.assertIn('Application Development Intern', full_text)
        self.assertIn('application development, including application design', full_text)
        self.assertIn('3 Months', full_text)
        self.assertIn('15/10/2026', full_text)
        self.assertNotIn('[Candidate Name]', full_text)
        self.assertNotIn('[1 Month / 3 Months]', full_text)

    def test_03_scenario_data_analytics_1_month(self):
        """Scenario 3: Candidate in Data Analytics with 1 Month duration uses Data Analytics template."""
        emp_doc, output_path = generate_offer_letter_docx(self.app_data)
        self.assertTrue(os.path.exists(output_path))

        doc = docx.Document(output_path)
        full_text = "\n".join([p.text for p in doc.paragraphs])

        self.assertIn('Chetan Verma', full_text)
        self.assertIn(self.app_data.formatted_code, full_text)
        self.assertIn('Data Analytics Intern', full_text)
        self.assertIn('data collection, data cleaning, data preprocessing', full_text)
        self.assertIn('1 Month', full_text)
        self.assertIn('01/11/2026', full_text)
        self.assertNotIn('[Candidate Name]', full_text)
        self.assertNotIn('[1 Month / 3 Months]', full_text)

    def test_04_scenario_full_stack_3_months(self):
        """Scenario 4: Candidate in Full Stack Development with 3 Months duration uses Full Stack template."""
        emp_doc, output_path = generate_offer_letter_docx(self.app_fullstack)
        self.assertTrue(os.path.exists(output_path))

        doc = docx.Document(output_path)
        full_text = "\n".join([p.text for p in doc.paragraphs])

        self.assertIn('Divya Nair', full_text)
        self.assertIn(self.app_fullstack.formatted_code, full_text)
        self.assertIn('Full Stack Development Intern', full_text)
        self.assertIn('full stack development, including frontend development, backend development', full_text)
        self.assertIn('3 Months', full_text)
        self.assertNotIn('[Candidate Name]', full_text)
        self.assertNotIn('[1 Month / 3 Months]', full_text)

    def test_05_scenario_unconfigured_domain_blocks_generation(self):
        """Scenario 5: Unconfigured domain raises exact required error message."""
        with self.assertRaises(OfferLetterTemplateNotFoundError) as ctx:
            generate_offer_letter_docx(self.app_unconfigured)
        
        expected_msg = "No active Offer Letter template is configured for Culinary Arts — 1 Month. Please upload or activate the appropriate template."
        self.assertEqual(str(ctx.exception), expected_msg)

    def test_06_duplicate_protection_and_deactivation(self):
        """Scenario 6: Uploading a new template for the same Job/Domain + Duration deactivates previous active template."""
        self.login_admin()

        # Get initial active AI & ML template
        initial_tmpl = DocumentTemplate.query.filter_by(
            template_type='offer_letter',
            job_domain='AI & ML',
            is_active=True
        ).first()
        self.assertIsNotNone(initial_tmpl)
        self.assertTrue(initial_tmpl.is_active)

        # Upload a replacement template for AI & ML (Both)
        docx_bytes = io.BytesIO()
        test_doc = docx.Document()
        test_doc.add_paragraph("AI & ML Custom Replacement Template")
        test_doc.add_paragraph("Dear [Candidate Name], Ref: [Reference Number], Duration: [1 Month / 3 Months]")
        test_doc.save(docx_bytes)
        docx_bytes.seek(0)

        res = self.client.post('/admin/templates/document/offer-letter/upload', data={
            'document_type': 'offer_letter',
            'job_domain': 'AI & ML',
            'duration': 'Both',
            'template_name': 'AI & ML V2 Custom Template',
            'template_file': (docx_bytes, 'ai_ml_v2.docx')
        }, content_type='multipart/form-data', follow_redirects=True)
        self.assertEqual(res.status_code, 200)

        # Refresh initial template from DB
        db.session.refresh(initial_tmpl)
        self.assertFalse(initial_tmpl.is_active, "Old template must be deactivated")

        # New template should be active
        new_tmpl = DocumentTemplate.query.filter_by(
            name='AI & ML V2 Custom Template',
            job_domain='AI & ML'
        ).first()
        self.assertIsNotNone(new_tmpl)
        self.assertTrue(new_tmpl.is_active, "New replacement template must be active")

    def test_07_historical_document_deletion_protection(self):
        """Scenario 7: Attempting to delete a template referenced by existing generated documents is blocked."""
        self.login_admin()

        # Generate an offer letter so template is referenced
        emp_doc, _ = generate_offer_letter_docx(self.app_aiml)
        tmpl_id = emp_doc.template_id
        self.assertIsNotNone(tmpl_id)

        # Attempt to delete referenced template
        res = self.client.post(f'/admin/templates/document/{tmpl_id}/delete', follow_redirects=True)
        self.assertEqual(res.status_code, 200)

        # Verify template still exists in DB
        tmpl = DocumentTemplate.query.get(tmpl_id)
        self.assertIsNotNone(tmpl, "Referenced template must NOT be deleted from DB")

    def test_08_admin_ui_endpoints(self):
        """Test admin template preview-info, download, and activate/deactivate routes."""
        self.login_admin()

        tmpl = DocumentTemplate.query.filter_by(template_type='offer_letter', is_active=True).first()
        self.assertIsNotNone(tmpl)

        # 1. Preview Info AJAX endpoint
        res = self.client.get(f'/admin/templates/document/{tmpl.id}/preview-info')
        self.assertEqual(res.status_code, 200)
        json_data = res.get_json()
        self.assertEqual(json_data['id'], tmpl.id)
        self.assertIn('job_domain', json_data)
        self.assertIn('detected_placeholders', json_data)
        self.assertTrue(isinstance(json_data['detected_placeholders'], list))

        # 2. Download endpoint
        res_dl = self.client.get(f'/admin/templates/document/{tmpl.id}/download')
        self.assertEqual(res_dl.status_code, 200)
        self.assertIn('application/vnd.openxmlformats-officedocument.wordprocessingml.document', res_dl.content_type)

        # 3. Deactivate endpoint
        res_deact = self.client.post(f'/admin/templates/document/{tmpl.id}/deactivate', follow_redirects=True)
        self.assertEqual(res_deact.status_code, 200)
        db.session.refresh(tmpl)
        self.assertFalse(tmpl.is_active)

        # 4. Activate endpoint
        res_act = self.client.post(f'/admin/templates/document/{tmpl.id}/activate', follow_redirects=True)
        self.assertEqual(res_act.status_code, 200)
        db.session.refresh(tmpl)
        self.assertTrue(tmpl.is_active)


if __name__ == '__main__':
    unittest.main()
